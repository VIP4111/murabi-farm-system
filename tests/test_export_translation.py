"""بند إصلاح (فحص عميق — طلبك: "ابدا بند التصدير") — ملفات تصدير Excel
(السجل المالي الكامل، تحليل نقطة التعادل، سجل التقريع) كانت تُبنى بأعمدة
وقيَم عربي بحت دايماً بغض النظر عن لغة المستخدم الحالي (`current_user.language`)،
بخلاف باقي الشاشات اللي عولجت هذا الفحص العميق. مستندات رسمية موجّهة
لطرف خارجي ثابت (فاتورة بيع للمشتري، مسير راتب للعامل، بروفايل دفعة
للمشتري المحتمل) مؤجلة عمداً خارج نطاق هذا البند — تبقى عربي بتصميمها،
هذا يخص فقط شاشات التصدير التحليلية اللي يستخدمها صاحب/فريق المزرعة
نفسه."""
import io

import openpyxl
from flask_babel import force_locale

from app.extensions import db
from app.models import Role, User, Finance, Mating
from datetime import date


def _english_owner():
    role = Role.query.filter_by(name="owner").first()
    u = User(name="EO", phone="0500099911", role_id=role.id, language="en")
    u.set_password("pass1234")
    db.session.add(u)
    db.session.commit()
    return u


def _login(client, user):
    client.post("/login", data={"phone": user.phone, "password": "pass1234"})


def test_finance_export_columns_translate_to_english(app, client):
    owner = _english_owner()
    db.session.add(Finance(date=date.today(), operation_type="sale", category="بيع رأس",
                            item="ت1", amount=100))
    db.session.commit()
    _login(client, owner)
    resp = client.get("/finance/export")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "التاريخ" not in header
    assert "الفئة" not in header
    row2 = [c.value for c in list(ws.iter_rows())[1]]
    # النوع (sale) والفئة (بيع رأس) مترجمين بصف البيانات نفسه
    assert "بيع" not in row2
    assert "بيع رأس" not in row2


def test_break_even_export_columns_translate_to_english(app, client):
    owner = _english_owner()
    _login(client, owner)
    resp = client.get("/finance/break-even-report/export")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "رقم الرأس" not in header
    assert "مصدر التقدير" not in header


def test_matings_export_columns_and_barn_translate_to_english(app, client):
    from app.models import Barn
    from tests.factories import make_animal
    owner = _english_owner()
    barn = Barn(barn_no="1", barn_name="حظيرة أ", barn_name_en="Barn A")
    db.session.add(barn)
    female = make_animal(animal_no="F-01", gender="أنثى")
    db.session.add(female)
    db.session.commit()
    db.session.add(Mating(date=date.today(), female_id=female.id, barn_id=barn.id, male_note="فحل خارجي"))
    db.session.commit()
    _login(client, owner)
    resp = client.get("/repro/matings/export")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "الحظيرة" not in header
    row2 = [c.value for c in list(ws.iter_rows())[1]]
    assert "حظيرة أ" not in row2
    assert "Barn A" in row2


def test_finance_export_stays_arabic_for_arabic_user(app, client, owner):
    # يثبت إن الترجمة تصير وقت العرض بس — نفس المستخدم العربي يشوف
    # الأعمدة عربي كالمعتاد (صفر تغيير على السلوك القديم).
    db.session.add(Finance(date=date.today(), operation_type="sale", category="بيع رأس", amount=50))
    db.session.commit()
    _login(client, owner)
    resp = client.get("/finance/export")
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "التاريخ" in header
